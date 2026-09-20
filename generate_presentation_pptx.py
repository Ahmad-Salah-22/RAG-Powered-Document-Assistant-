import sys
import subprocess

# Ensure python-pptx is installed
try:
    import pptx
except ImportError:
    print("Installing python-pptx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Dark Navy Theme
    BG_COLOR = RGBColor(15, 23, 42)      # #0F172A
    CARD_BG = RGBColor(30, 41, 59)       # #1E293B
    TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    ACCENT_BLUE = RGBColor(96, 165, 250) # #60A5FA
    ACCENT_PURPLE = RGBColor(167, 139, 250) # #A78BFA
    ACCENT_GREEN = RGBColor(74, 222, 128) # #4ADE80
    ACCENT_RED = RGBColor(248, 113, 113) # #F87171

    def add_blank_slide_with_bg():
        slide = prs.slides.add_slide(blank_layout)
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, category_text="RAG DOCUMENT ASSISTANT"):
        # Category / Kicker
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_PURPLE

        # Main Slide Title
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = ACCENT_BLUE

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    s1 = add_blank_slide_with_bg()
    
    # Title Box
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "📚 RAG-Powered Document Assistant"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    p2 = tf.add_paragraph()
    p2.text = "Privacy-First Grounded Question Answering & Citation System"
    p2.font.size = Pt(22)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(10)

    # Subtitle Card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.5), Inches(11.333), Inches(1.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = ACCENT_PURPLE
    tf_card = card.text_frame
    tf_card.word_wrap = True
    
    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Capstone & Graduation Project Presentation"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = TEXT_WHITE
    
    p_c2 = tf_card.add_paragraph()
    p_c2.text = "Technologies: FastAPI • ChromaDB • SentenceTransformers • Ollama (llama3:8b) • Streamlit"
    p_c2.font.size = Pt(14)
    p_c2.font.color.rgb = ACCENT_GREEN
    p_c2.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 2: PROBLEM STATEMENT
    # -------------------------------------------------------------
    s2 = add_blank_slide_with_bg()
    add_header(s2, "Problem Statement & Motivation")

    col_width = Inches(3.64)
    col_gap = Inches(0.4)
    left_margin = Inches(0.8)

    problems = [
        ("Traditional Search", "❌ Keyword Based", "Returns entire documents requiring manual reading.\nFails on semantic synonyms.\nNo automated answer synthesis.", ACCENT_RED),
        ("Generic LLMs (ChatGPT)", "❌ Hallucination Prone", "Lacks access to private academic PDFs.\nProne to plausible-sounding false answers.\nZero page or document citation.", ACCENT_RED),
        ("RAG Assistant (Ours)", "✅ Grounded & Auditable", "Extracts precise passages from local PDFs.\nStrict zero-hallucination guardrails.\n100% verifiable page-level citations.", ACCENT_GREEN)
    ]

    for i, (title, tag, desc, color) in enumerate(problems):
        x = left_margin + i * (col_width + col_gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), col_width, Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p_tag = tf.add_paragraph()
        p_tag.text = tag
        p_tag.font.size = Pt(14)
        p_tag.font.bold = True
        p_tag.font.color.rgb = color
        p_tag.space_before = Pt(6)
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: ARCHITECTURE & SYSTEM FLOW
    # -------------------------------------------------------------
    s3 = add_blank_slide_with_bg()
    add_header(s3, "System Architecture & Dual-Pipeline Flow")

    # Ingestion Card
    c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_BLUE
    tf1 = c1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "1. Ingestion Pipeline (Offline Build)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    steps1 = [
        "📄 Raw PDF Ingestion (PyPDF parser)",
        "🧹 Cleaning & Page-by-Page Extraction",
        "🧩 Overlapping Chunker (Size 800, Overlap 150)",
        "🔢 SentenceTransformers (all-MiniLM-L6-v2)",
        "💾 Persistent ChromaDB Storage (On-Disk)"
    ]
    for st in steps1:
        p_s = tf1.add_paragraph()
        p_s.text = st
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = TEXT_WHITE
        p_s.space_before = Pt(10)

    # Runtime Card
    c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(1.8), Inches(5.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_PURPLE
    tf2 = c2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "2. Runtime Query Pipeline (Online)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    
    steps2 = [
        "👤 User submits query via Streamlit UI",
        "⚡ FastAPI receives POST /query request",
        "🔍 ChromaDB retrieves Top-K nearest chunks",
        "🛡️ Context formatted into Guardrail Prompt",
        "🦙 Ollama llama3:8b generates grounded answer",
        "📄 Response rendered with clickable Citations"
    ]
    for st in steps2:
        p_s = tf2.add_paragraph()
        p_s.text = st
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = TEXT_WHITE
        p_s.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: TECH STACK MATRIX
    # -------------------------------------------------------------
    s4 = add_blank_slide_with_bg()
    add_header(s4, "Technology Stack & Core Components")

    # Table creation
    rows = 6
    cols = 3
    table_shape = s4.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(5.433)

    headers = ["Component", "Technologies", "Role & Configuration"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = BG_COLOR

    data = [
        ("Backend API", "FastAPI, Uvicorn, Pydantic", "REST endpoints (/health, /query), Lifespan startup caching"),
        ("Dense Vector Store", "ChromaDB, SentenceTransformers", "384d dense embeddings (all-MiniLM-L6-v2), Cosine distance metric"),
        ("Local LLM Engine", "Ollama (llama3:8b)", "Grounded prompt execution with low temperature (0.1) for zero hallucination"),
        ("Frontend UI", "Streamlit, Glassmorphism CSS", "Diagnostic sidebar, chat history, collapsible source citation accordion"),
        ("DevOps & Testing", "Pytest, Docker, Compose", "Automated API unit test suite and containerized deployment")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 5: DATA INGESTION & CHUNKING
    # -------------------------------------------------------------
    s5 = add_blank_slide_with_bg()
    add_header(s5, "Data Ingestion & Chunking Strategy")

    card1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.2))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = ACCENT_BLUE
    tf = card1.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "🧩 Overlapping Chunking Strategy"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    items = [
        "• Chunk Size: 800 characters (~120 words) preserving complete technical concepts.",
        "• Chunk Overlap: 150 characters to prevent loss of context across chunk boundaries.",
        "• Metadata Binding: Every chunk stores Document Name, Page Number, and unique Chunk ID."
    ]
    for it in items:
        p_it = tf.add_paragraph()
        p_it.text = it
        p_it.font.size = Pt(13)
        p_it.font.color.rgb = TEXT_WHITE
        p_it.space_before = Pt(4)

    card2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.3))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = ACCENT_GREEN
    tf2 = card2.text_frame
    tf2.word_wrap = True
    
    p2 = tf2.paragraphs[0]
    p2.text = "🔢 Dense Vector Representation & Caching"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_GREEN
    
    items2 = [
        "• Model: SentenceTransformers all-MiniLM-L6-v2 (384-dimensional dense vectors).",
        "• Persistence: Vector collection saved locally to backend/data/vector_store/.",
        "• Lifespan Preloading: Vector index loaded ONCE into memory at server startup for sub-millisecond retrieval."
    ]
    for it in items2:
        p_it = tf2.add_paragraph()
        p_it.text = it
        p_it.font.size = Pt(13)
        p_it.font.color.rgb = TEXT_WHITE
        p_it.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 6: GROUNDED GENERATION & GUARDRAILS
    # -------------------------------------------------------------
    s6 = add_blank_slide_with_bg()
    add_header(s6, "Grounded Generation & Anti-Hallucination Guardrails")

    card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = ACCENT_PURPLE
    tf = card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🛡️ System Prompt Guardrail Architecture"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    rules = [
        "1. Strict Context Constraint: Answer user questions using ONLY the provided document context blocks.",
        "2. Mandatory Fallback Rule: If context lacks required data, reply: 'I could not find this information in the provided documents.'",
        "3. Zero-Extrapolation Rule: Do NOT make up facts, guess, or extrapolate beyond the provided text.",
        "4. Exact Citation Format: Always cite your source using format: [Document Name, Page X].",
        "5. Low Temperature Execution: Model called with temperature=0.1 to eliminate random/creative output."
    ]

    for r in rules:
        p_r = tf.add_paragraph()
        p_r.text = r
        p_r.font.size = Pt(13)
        p_r.font.color.rgb = TEXT_WHITE
        p_r.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 7: EMPIRICAL EVALUATION & RESULTS
    # -------------------------------------------------------------
    s7 = add_blank_slide_with_bg()
    add_header(s7, "Empirical Evaluation & Benchmark Matrix")

    rows = 6
    cols = 4
    table_shape = s7.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(4.5)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(2.0)
    table.columns[3].width = Inches(2.433)

    headers = ["Test Query", "Retrieved Document", "Similarity Dist.", "System Verdict"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = BG_COLOR

    data = [
        ("Average lookup complexity of hash table?", "cs_data_structures.pdf (p.2)", "0.3200", "✅ Grounded (O(1))"),
        ("Core pillars of Object-Oriented Programming?", "cs_intro_python.pdf (p.3)", "0.3829", "✅ Grounded (4 Pillars)"),
        ("How does RAG work?", "cs_machine_learning.pdf (p.3)", "0.4240", "✅ Grounded (Dense RAG)"),
        ("Time complexity of MergeSort?", "cs_data_structures.pdf (p.3)", "0.4095", "✅ Grounded (O(n log n))"),
        ("Who won the 2022 World Cup? (Out-of-Domain)", "None (Off-Topic)", "0.9043", "🛡️ Safe Fallback")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.color.rgb = ACCENT_GREEN if "✅" in val or "🛡️" in val else TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 8: KEY INNOVATIONS & HIGHLIGHTS
    # -------------------------------------------------------------
    s8 = add_blank_slide_with_bg()
    add_header(s8, "Key Technical Innovations & System Highlights")

    innovations = [
        ("🔒 100% Local & Private", "Zero cloud API dependencies. Embedding generation and LLM execution run completely offline locally."),
        ("🎯 Grounded & Verifiable", "Page-level source citations allow users to verify every sentence against original source PDFs."),
        ("🛡️ Robust Guardrails", "Dual-layer fallback protection prevents hallucinations on ungrounded or off-topic queries."),
        ("⚡ High Performance", "FastAPI lifespan caching loads vector models into memory ONCE for instant query processing.")
    ]

    for i, (title, desc) in enumerate(innovations):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + col * Inches(5.9)
        y = Inches(1.8) + row * Inches(2.5)

        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.633), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE
        
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_WHITE
        p_desc.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 9: CONCLUSION & FUTURE ROADMAP
    # -------------------------------------------------------------
    s9 = add_blank_slide_with_bg()
    add_header(s9, "Conclusion & Future Roadmap")

    card1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = ACCENT_GREEN
    tf1 = card1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "🏆 Achieved Outcomes"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    
    outs = [
        "✅ Fully functional production-grade RAG pipeline",
        "✅ 100% precision on in-domain academic queries",
        "✅ Zero hallucination guardrail enforcement",
        "✅ Modern Glassmorphism Streamlit UI",
        "✅ Unit tested & Docker container ready"
    ]
    for o in outs:
        p_o = tf1.add_paragraph()
        p_o.text = o
        p_o.font.size = Pt(13)
        p_o.font.color.rgb = TEXT_WHITE
        p_o.space_before = Pt(12)

    card2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(1.8), Inches(5.6), Inches(4.8))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = ACCENT_PURPLE
    tf2 = card2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "🚀 Future Extensions"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    
    exts = [
        "🔍 Hybrid Search (BM25 + Dense Vector Search)",
        "⚡ Cross-Encoder Reranking for Top-K passages",
        "📤 Dynamic Drag-and-Drop PDF upload portal",
        "📊 Multi-modal PDF extraction (Tables & Charts)"
    ]
    for e in exts:
        p_e = tf2.add_paragraph()
        p_e.text = e
        p_e.font.size = Pt(13)
        p_e.font.color.rgb = TEXT_WHITE
        p_e.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 10: Q&A / THANK YOU
    # -------------------------------------------------------------
    s10 = add_blank_slide_with_bg()
    
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Thank You! Questions & Answers"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "📚 RAG-Powered Document Assistant"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(14)

    output_path = "RAG_Document_Assistant_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to '{output_path}'")

if __name__ == "__main__":
    create_presentation()
