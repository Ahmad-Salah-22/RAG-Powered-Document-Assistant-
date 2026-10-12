import os
import io
import glob
from typing import List, Dict, Any, Tuple
import pypdf
import chromadb
from sentence_transformers import SentenceTransformer
from app.utils.logging_config import logger
from app.utils.text_enhancer import clean_and_enhance_text
from app.core.config import settings


def extract_pdf_pages(pdf_path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Extracts text page by page from a PDF file using PyPDF.

    Returns:
        Tuple of (list of page dicts, document stats dict)
    """
    doc_name = os.path.basename(pdf_path)
    pages = []
    stats = {
        "document": doc_name,
        "page_count": 0,
        "char_count": 0,
        "failed_pages": 0,
        "requires_ocr": False
    }

    try:
        reader = pypdf.PdfReader(pdf_path)
        stats["page_count"] = len(reader.pages)

        for page_num, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
                cleaned_text = clean_and_enhance_text(" ".join(text.split()))
                char_len = len(cleaned_text)

                if char_len < 20:  # Potential image-only / OCR page
                    stats["requires_ocr"] = True

                stats["char_count"] += char_len
                pages.append({
                    "document": doc_name,
                    "page": page_num,
                    "text": cleaned_text
                })
            except Exception as e:
                logger.error(f"Error extracting page {page_num} in '{doc_name}': {e}")
                stats["failed_pages"] += 1

    except Exception as e:
        logger.error(f"Failed to read PDF '{doc_name}': {e}")

    return pages, stats


import zipfile
import xml.etree.ElementTree as ET
import csv


def _extract_docx(file_bytes: bytes) -> str:
    """Extracts paragraph text from Word .docx file via package XML."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            if "word/document.xml" not in zf.namelist():
                return ""
            xml_content = zf.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for p in tree.findall(".//w:p", ns):
                texts = [node.text for node in p.findall(".//w:t", ns) if node.text]
                if texts:
                    paragraphs.append("".join(texts).strip())
            return "\n\n".join(paragraphs)
    except Exception as e:
        logger.error(f"Error parsing docx file: {e}")
        return ""


def _extract_pptx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts slide text from PowerPoint .pptx file."""
    pages = []
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = "".join([r.text for r in p.runs if r.text]).strip()
                        if txt:
                            slide_texts.append(txt)
            cleaned = " ".join(" ".join(slide_texts).split())
            pages.append({"page": slide_idx, "text": cleaned})
    except Exception as e:
        logger.error(f"Error parsing pptx presentation: {e}")
    return pages


def _extract_csv(file_bytes: bytes) -> str:
    """Extracts structured rows from CSV data."""
    try:
        text_content = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text_content))
        rows = list(reader)
        if not rows:
            return ""
        header = rows[0]
        lines = []
        for row_idx, r in enumerate(rows[1:], start=1):
            row_desc = ", ".join([f"{h.strip()}: {v.strip()}" for h, v in zip(header, r) if v.strip()])
            if row_desc:
                lines.append(f"Record {row_idx} ({row_desc})")
        return ". ".join(lines)
    except Exception as e:
        logger.error(f"Error parsing CSV file: {e}")
        return ""


def _extract_xlsx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts sheets from Excel spreadsheet."""
    pages = []
    try:
        import pandas as pd
        excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
        for sheet_idx, sheet_name in enumerate(excel_file.sheet_names, start=1):
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            sheet_text = f"Sheet: {sheet_name}. " + df.to_string(index=False)
            cleaned = " ".join(sheet_text.split())
            pages.append({"page": sheet_idx, "text": cleaned})
    except Exception as e:
        logger.error(f"Error parsing xlsx workbook: {e}")
    return pages


def extract_from_bytes(file_bytes: bytes, filename: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Extracts text and page metadata from in-memory file bytes.
    Supports PDF, DOCX, PPTX, CSV, XLSX, TXT, and Markdown files.
    """
    ext = os.path.splitext(filename)[1].lower()
    pages = []
    stats = {
        "document": filename,
        "page_count": 0,
        "char_count": 0,
        "failed_pages": 0,
        "requires_ocr": False,
        "file_type": ext.lstrip(".") or "unknown",
        "summary_preview": ""
    }

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            stats["page_count"] = len(reader.pages)
            for page_num, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                    cleaned_text = clean_and_enhance_text(" ".join(text.split()))
                    char_len = len(cleaned_text)
                    if char_len < 20:
                        stats["requires_ocr"] = True
                    stats["char_count"] += char_len
                    pages.append({
                        "document": filename,
                        "page": page_num,
                        "text": cleaned_text
                    })
                except Exception as e:
                    logger.error(f"Error reading page {page_num} of {filename}: {e}")
                    stats["failed_pages"] += 1
        except Exception as e:
            logger.error(f"Failed to parse PDF bytes for '{filename}': {e}")

    elif ext == ".docx":
        extracted_text = _extract_docx(file_bytes)
        cleaned_text = " ".join(extracted_text.split())
        stats["page_count"] = 1
        stats["char_count"] = len(cleaned_text)
        pages.append({
            "document": filename,
            "page": 1,
            "text": cleaned_text
        })

    elif ext == ".pptx":
        slide_pages = _extract_pptx(file_bytes)
        stats["page_count"] = len(slide_pages)
        for sp in slide_pages:
            stats["char_count"] += len(sp["text"])
            pages.append({
                "document": filename,
                "page": sp["page"],
                "text": sp["text"]
            })
        if not pages:
            pages.append({"document": filename, "page": 1, "text": ""})

    elif ext == ".csv":
        csv_text = _extract_csv(file_bytes)
        cleaned_text = " ".join(csv_text.split())
        stats["page_count"] = 1
        stats["char_count"] = len(cleaned_text)
        pages.append({
            "document": filename,
            "page": 1,
            "text": cleaned_text
        })

    elif ext in [".xlsx", ".xls"]:
        sheet_pages = _extract_xlsx(file_bytes)
        stats["page_count"] = len(sheet_pages)
        for sp in sheet_pages:
            stats["char_count"] += len(sp["text"])
            pages.append({
                "document": filename,
                "page": sp["page"],
                "text": sp["text"]
            })
        if not pages:
            pages.append({"document": filename, "page": 1, "text": ""})

    elif ext in [".txt", ".md", ".markdown"]:
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            cleaned_text = " ".join(text.split())
            stats["page_count"] = 1
            stats["char_count"] = len(cleaned_text)
            pages.append({
                "document": filename,
                "page": 1,
                "text": cleaned_text
            })
        except Exception as e:
            logger.error(f"Failed to decode text file '{filename}': {e}")
    else:
        text = file_bytes.decode("utf-8", errors="ignore")
        cleaned_text = " ".join(text.split())
        stats["page_count"] = 1
        stats["char_count"] = len(cleaned_text)
        pages.append({
            "document": filename,
            "page": 1,
            "text": cleaned_text
        })

    # Generate document abstract / preview snippet with clean, pleasant formatting
    all_text = " ".join(p["text"] for p in pages if p.get("text"))
    if all_text:
        cleaned_preview = clean_and_enhance_text(all_text[:600])
        stats["summary_preview"] = cleaned_preview[:280] + ("..." if len(cleaned_preview) > 280 else "")

    return pages, stats


def chunk_text(
    pages: List[Dict[str, Any]],
    chunk_size: int = settings.CHUNK_SIZE,
    chunk_overlap: int = settings.CHUNK_OVERLAP
) -> List[Dict[str, Any]]:
    """
    Splits page text into overlapping chunks while preserving page metadata.
    """
    chunks = []
    global_chunk_counter = 0

    for page_info in pages:
        text = page_info["text"]
        doc_name = page_info["document"]
        page_num = page_info["page"]

        if not text:
            continue

        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk_str = text[start:end].strip()

            if chunk_str:
                global_chunk_counter += 1
                chunk_id = f"{doc_name}_p{page_num}_c{global_chunk_counter}"
                chunks.append({
                    "id": chunk_id,
                    "text": chunk_str,
                    "metadata": {
                        "document": doc_name,
                        "page": page_num,
                        "chunk_id": chunk_id,
                        "char_length": len(chunk_str)
                    }
                })

            if end == text_len:
                break
            start += (chunk_size - chunk_overlap)

    return chunks


def index_uploaded_file(
    file_bytes: bytes,
    filename: str,
    collection,
    embedding_model: SentenceTransformer,
    raw_data_dir: str = "./data/raw"
) -> Dict[str, Any]:
    """
    Takes an uploaded file's bytes, saves a copy in raw_data_dir, extracts, chunks,
    generates embeddings, and appends to the active ChromaDB collection.
    """
    os.makedirs(raw_data_dir, exist_ok=True)
    saved_path = os.path.join(raw_data_dir, filename)
    with open(saved_path, "wb") as f:
        f.write(file_bytes)
    logger.info(f"Saved uploaded file to disk at '{saved_path}'")

    # If document already existed in collection, remove old chunks first
    try:
        existing = collection.get(where={"document": filename})
        if existing and existing.get("ids"):
            collection.delete(ids=existing["ids"])
            logger.info(f"Cleaned up {len(existing['ids'])} existing chunks for overwritten document '{filename}'")
    except Exception as e:
        logger.warning(f"Could not purge existing chunks for '{filename}': {e}")

    pages, stats = extract_from_bytes(file_bytes, filename)
    chunks = chunk_text(pages)

    if not chunks:
        logger.warning(f"No chunks created from uploaded file '{filename}'")
        return {
            "document_name": filename,
            "pages": len(pages),
            "chunks_indexed": 0,
            "char_count": stats.get("char_count", 0)
        }

    chunk_texts = [c["text"] for c in chunks]
    embeddings = embedding_model.encode(chunk_texts, show_progress_bar=False).tolist()

    ids = [c["id"] for c in chunks]
    metadatas = [c["metadata"] for c in chunks]

    batch_size = 100
    for i in range(0, len(chunks), batch_size):
        collection.add(
            ids=ids[i:i+batch_size],
            embeddings=embeddings[i:i+batch_size],
            documents=chunk_texts[i:i+batch_size],
            metadatas=metadatas[i:i+batch_size]
        )

    logger.info(f"Successfully indexed {len(chunks)} chunks for uploaded file '{filename}'.")
    return {
        "document_name": filename,
        "pages": len(pages),
        "chunks_indexed": len(chunks),
        "char_count": stats.get("char_count", 0),
        "file_type": stats.get("file_type", "pdf"),
        "requires_ocr": stats.get("requires_ocr", False),
        "summary_preview": stats.get("summary_preview", "")
    }


def build_and_persist_vector_store(
    raw_data_dir: str = "./data/raw",
    chroma_path: str = settings.CHROMA_PATH,
    collection_name: str = settings.COLLECTION_NAME,
    embedding_model_name: str = settings.EMBEDDING_MODEL_NAME
) -> Dict[str, Any]:
    """
    Full pipeline to load raw documents (PDF, DOCX, PPTX, CSV, XLSX, TXT, MD), chunk, compute embeddings, and persist to ChromaDB.
    """
    logger.info("=== Starting Data Ingestion Pipeline ===")
    supported_extensions = ["*.pdf", "*.txt", "*.md", "*.markdown", "*.docx", "*.pptx", "*.csv", "*.xlsx"]
    doc_files = []
    for ext in supported_extensions:
        doc_files.extend(glob.glob(os.path.join(raw_data_dir, ext)))
    
    if not doc_files:
        logger.warning(f"No documents found in raw data directory '{raw_data_dir}'.")
        return {"documents_processed": 0, "total_chunks": 0}

    all_pages = []
    all_doc_stats = []

    for file_path in doc_files:
        logger.info(f"Processing document: '{file_path}'")
        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            b = f.read()
        pages, stats = extract_from_bytes(b, filename)
        all_pages.extend(pages)
        all_doc_stats.append(stats)

    chunks = chunk_text(all_pages)
    logger.info(f"Extracted {len(all_pages)} total pages across {len(doc_files)} documents.")
    logger.info(f"Generated {len(chunks)} text chunks (size={settings.CHUNK_SIZE}, overlap={settings.CHUNK_OVERLAP}).")

    if not chunks:
        logger.warning("No text chunks generated.")
        return {"documents_processed": len(doc_files), "total_chunks": 0}

    # Load embedding model
    model = SentenceTransformer(embedding_model_name)
    chunk_texts = [c["text"] for c in chunks]

    logger.info(f"Generating embeddings using '{embedding_model_name}'...")
    embeddings = model.encode(chunk_texts, show_progress_bar=False).tolist()

    # ChromaDB persistent store initialization
    abs_chroma_path = os.path.abspath(chroma_path)
    os.makedirs(abs_chroma_path, exist_ok=True)

    client = chromadb.PersistentClient(path=abs_chroma_path)
    
    # Reset/recreate collection for fresh indexing
    try:
        client.delete_collection(name=collection_name)
        logger.info(f"Reset existing collection '{collection_name}'.")
    except Exception:
        pass

    collection = client.create_collection(
        name=collection_name,
        metadata={"hnsw:space": "cosine"}
    )

    ids = [c["id"] for c in chunks]
    metadatas = [c["metadata"] for c in chunks]

    # Add to ChromaDB in batches
    batch_size = 100
    for i in range(0, len(chunks), batch_size):
        collection.add(
            ids=ids[i:i+batch_size],
            embeddings=embeddings[i:i+batch_size],
            documents=chunk_texts[i:i+batch_size],
            metadatas=metadatas[i:i+batch_size]
        )

    logger.info(f"Successfully saved {collection.count()} vectors into persistent database at '{abs_chroma_path}'.")

    pipeline_summary = {
        "documents_processed": len(doc_files),
        "total_pages": sum(s["page_count"] for s in all_doc_stats),
        "total_chunks": len(chunks),
        "vector_store_path": abs_chroma_path,
        "doc_stats": all_doc_stats
    }

    logger.info("=== Ingestion Pipeline Completed Successfully ===")
    return pipeline_summary


if __name__ == "__main__":
    summary = build_and_persist_vector_store()
    print("Ingestion Summary:", summary)
